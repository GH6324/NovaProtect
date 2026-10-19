from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '宣发资料'
OUT.mkdir(exist_ok=True)
DOCX = OUT / 'NovaProtect产品宣发文案.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(19), Mm(18)
section.left_margin, section.right_margin = Mm(21), Mm(21)
section.header_distance, section.footer_distance = Mm(8), Mm(8)

def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for role in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rfonts.set(qn('w:' + role), 'Microsoft YaHei')
    for child in list(rpr):
        if child.tag == qn('w:color'):
            for key in ['themeColor', 'themeTint', 'themeShade']:
                child.attrib.pop(qn('w:' + key), None)

normal = doc.styles['Normal']
font(normal, 11)
normal.paragraph_format.line_spacing = 1.32
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 17), ('Heading 2', 12)]:
    style = doc.styles[name]
    font(style, size, name != 'Subtitle')
    style.paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 0)
    style.paragraph_format.space_after = Pt(7 if name == 'Heading 2' else 11)
    style.paragraph_format.keep_with_next = True
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('NovaProtect  |  ')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string('666666')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, bold_lead=None):
    para = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para

def h(text):
    doc.add_paragraph(text, 'Heading 2')

def page(title):
    doc.add_page_break()
    doc.add_paragraph(title, 'Heading 1')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + name)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top', '100'), ('bottom', '100'), ('left', '120'), ('right', '120')]:
        el = OxmlElement('w:' + name)
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for ri, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill = '153F60' if ri == 0 else ('F1F5F8' if ri % 2 == 0 else 'FFFFFF')
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), fill)
            cell._tc.get_or_add_tcPr().append(shading)
            for para in cell.paragraphs:
                if len(headers) > 2 and ci >= 2:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.2
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if ri == 0 else '000000')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

logo = doc.add_paragraph()
logo.paragraph_format.space_after = Pt(8)
logo.add_run().add_picture(str(ROOT / 'ui' / 'Assets' / 'NovaProtect.png'), width=Mm(14))
doc.add_paragraph('NovaProtect 产品宣发文案', 'Title')
doc.add_paragraph('面向 Windows 开发者的代码保护工作台', 'Subtitle')
p('NovaProtect 将目标分析、保护配置、关键区域管理和构建输出整合到统一工作台，帮助开发者为核心算法与业务逻辑配置保护措施。开发者可以根据目标程序的架构、兼容性和运行开销，选择适合项目的保护组合。')
p('产品可处理 C/C++、Rust、Go 等工具链生成的受支持原生 EXE，并为 C#、VB .NET EXE 提供加密封装。实际能力由文件格式、处理器架构和可处理的代码区域共同决定。')
h('产品价值')
p('关键逻辑按需保护。通过 NovaProtect SDK 指定重点区域，在受支持的原生 x64 程序中应用局部虚拟化及 VM 前置变异，增加理解关键代码所需的分析工作。', '关键逻辑按需保护。')
p('配置效果清晰可见。构建前展示请求项、生效项、跳过项与冲突原因；项目文件和历史快照保留配置与处理结果，便于持续维护和问题定位。', '配置效果清晰可见。')
h('软件构成')
table(['组成部分', '主要作用'], [
    ['可视化客户端', 'C# 与 WinUI 3 界面，提供项目配置、区域编辑、构建进度和日志。'],
    ['原生保护引擎', 'C++ 实现 PE 分析、区段处理、虚拟指令转换与输出构建。'],
    ['保护运行时', '包含 x86/x64 加载器、x64 解释器与 .NET 专用启动宿主。'],
    ['SDK 与项目管理', '标记关键区域，保存区域规则、项目配置和构建快照。'],
    ['官网与分发服务', '提供客户端下载、更新日志和后台发布入口，无需账户即可获取安装包。'],
], [36, 132])

page('加密与保护原理')
p('NovaProtect 采用分层处理流程。引擎先识别目标程序的结构与能力，再对符合条件的区段、SDK 区域及导入信息应用相应措施，最后生成保护输出与构建记录。')
h('原生区段保护')
p('针对符合条件的代码区段以及可选资源区段进行打包；在压缩能够缩小载荷时使用压缩，并在启用区段加密时采用 XOR 编码。程序启动后，由对应架构的加载器恢复载荷。该路径用于改变程序在磁盘上的直接表示形式。')
h('关键区域虚拟化与代码变异')
p('针对原生 x64 程序中由 NovaProtect SDK 标记、能够成功翻译的代码区域，将机器指令转换为内部虚拟指令，运行时交由解释器执行。虚拟指令采用映射表编码，Ultra 管线还可在转换前执行代码变异，并对适用字节码追加混淆。')
p('这使关键区域的分析需要进一步理解内部指令与执行流程。具体区域能否处理，以指令支持情况、区域空间和构建结果为准。')
h('VM 字节码组合加密')
p('生成的 VM 字节码采用 RC4、XOR 与字节加法组合处理，并结合区域随机参数及区域位置派生处理参数。运行时按对应顺序恢复字节码，再交给解释器执行。')
h('导入保护与运行时校验')
p('原生 x86/x64 导入保护保存导入信息，并由加载器在运行时重建导入地址表。根据目标与配置，还可执行载荷或代码区域完整性校验；原生 x64 进一步提供许可数据校验和启动前机器标识绑定。')
h('托管程序加密封装')
p('对受支持的 C#、VB .NET EXE，使用基于 XTEA 的载荷加密，由专用宿主完成解密、完整性校验与程序启动。当前方案保留原始 CLR 元数据和 IL 结构，采用封装方式接入原有运行环境。')

page('设计理念与适用范围')
p('NovaProtect 的设计着重于让开发者理解保护对象、配置作用和构建结果，使代码保护能够融入项目维护与版本发布流程。')
h('保护对象明确')
p('区域对象来自目标文件中的真实 NovaProtect SDK 标记。开发者可以结合代码价值、性能要求和兼容性，确定需要处理的关键区域，并为区域设置相应规则。')
h('配置与结果可核查')
p('兼容、均衡、高强度和 Ultra 等预设提供配置起点。构建计划进一步说明哪些请求实际生效、哪些被跳过以及是否存在冲突；历史快照记录项目配置与构建结果，便于后续复用和核查。')
h('产品设计参考')
p('NovaProtect 的产品设计主要参考 VMProtect Professional 与 Themida。VMProtect Professional 提供了项目与代码区域组织、属性配置、保护模式和构建日志方面的参考；Themida 提供了 SDK 宏标记关键区域、按模块组织保护能力方面的参考。')
p('这些参考体现在产品组织与操作流程中，不代表复现相关产品的全部功能、内部算法或保护强度。NovaProtect 围绕自身的 SDK、保护引擎和构建计划提供对应功能。')
h('当前能力概览')
table(['目标类型', '区段或载荷保护', 'SDK 区域 VM', '许可与机器绑定'], [
    ['原生 x86 EXE', '区段打包与编码', '暂不支持', '暂不支持'],
    ['原生 x64 EXE', '区段打包与编码', '支持适用区域', '支持'],
    ['C# 与 VB .NET EXE', '加密封装', '暂不支持', '暂不支持'],
    ['ARM64', '只读识别', '暂不支持', '暂不支持'],
], [36, 53, 39, 40])
p('原生导入保护覆盖受支持的 x86/x64 目标；.NET 当前提供加密封装，尚未实现 IL 虚拟化；ARM64 暂不生成保护输出。实际可用项目以客户端构建计划和目标程序验证结果为准。')

page('可直接发布的宣传文案')
h('公众号与产品发布介绍')
p('NovaProtect 面向 Windows 开发者，将代码保护与日常发布流程连接起来。围绕商业软件中的核心算法和关键业务逻辑，产品提供区段保护、原生 x64 局部虚拟化、VM 字节码组合加密、导入保护与运行时校验，并为 C#、VB .NET EXE 提供加密封装。')
p('开发者可以通过 SDK 标记需要重点处理的代码区域，再根据项目的兼容性、运行开销和保护需求选择配置。可视化客户端集中呈现目标信息、区域规则、构建进度与日志，帮助开发者了解每项设置在当前程序中的作用。')
p('NovaProtect 将配置透明和结果追溯作为产品设计重点。构建前展示生效项、跳过项与冲突原因；构建后保留项目配置和历史快照，方便持续维护、复用配置及定位兼容性问题。')
p('产品设计参考了 VMProtect Professional 的项目与区域管理思路，以及 Themida 的 SDK 区域标记与模块组织方式，并在自身工作台中形成统一操作流程。对于需要为 Windows 软件配置代码保护的开发者，NovaProtect 提供了从目标分析到输出核查的集中入口。')
h('社群发布短文案')
p('NovaProtect 代码保护工作台，面向 Windows 开发者提供区段保护、原生 x64 SDK 区域虚拟化、字节码加密、导入保护与运行时校验，并支持 C#、VB .NET EXE 加密封装。可视化配置、生效结果提示与构建快照，让关键代码保护更清晰，让版本维护更有据可查。')
h('官网简介')
p('NovaProtect 是面向 Windows EXE 的代码保护工作台，整合目标分析、SDK 区域管理、分层保护与构建记录。针对不同架构和文件格式提供相应能力，帮助开发者按需配置关键逻辑保护，并追溯实际构建结果。')
h('产品信息')
p('产品名称  NovaProtect\n产品定位  Windows EXE 代码保护工作台\n面向用户  Windows 软件开发者与软件发布团队\n产品官网  nova.fbi.plus')

doc.core_properties.title = 'NovaProtect 产品宣发文案'
doc.core_properties.subject = '产品介绍 加密保护原理 软件构成 设计理念 设计参考 宣传文案'
doc.core_properties.author = 'NovaProtect'
doc.core_properties.keywords = 'NovaProtect Windows EXE 代码保护 产品介绍'
doc.core_properties.comments = ''
doc.save(DOCX)
print(DOCX)

# Retain an editable text companion for publishing and later updates.
parts = []
for child in doc.element.body:
    if child.tag == qn('w:p'):
        txt = ''.join(child.itertext())
        # python-docx Paragraph.text avoids duplicate text from OOXML traversal.
        from docx.text.paragraph import Paragraph
        para = Paragraph(child, doc)
        if para.text.strip():
            if para.style.name in ['Title', 'Heading 1', 'Heading 2']:
                parts.append('**' + para.text + '**')
            else:
                parts.append(para.text)
    elif child.tag == qn('w:tbl'):
        from docx.table import Table
        t = Table(child, doc)
        rows = [[c.text.replace('\n', ' ') for c in r.cells] for r in t.rows]
        parts.append('\n'.join(['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * len(rows[0])) + ' |'] + ['| ' + ' | '.join(r) + ' |' for r in rows[1:]]))
(OUT / 'NovaProtect产品宣发文案.md').write_text('\n\n'.join(parts) + '\n', encoding='utf-8')
print('Paragraphs:', len(doc.paragraphs), 'Tables:', len(doc.tables))
